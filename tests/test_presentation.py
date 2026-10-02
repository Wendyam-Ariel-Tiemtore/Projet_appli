"""Présentations orales et choix de méthodes de l'auteur."""

import re
import zipfile

import pytest

from analyste.writing.presentation import GENRES, PresentationSpec

TIRETS = re.compile(r"[—–]")


def test_options_valides_et_valeurs_par_defaut_du_genre():
    p = PresentationSpec.from_dict({"active": True, "genre": "restitution", "theme": "inconnu", "duree": "999",
                                    "contenus": ["bivarie", "piratage"]})
    assert p.genre == "restitution" and p.theme == "ardoise"
    assert p.duree == GENRES["restitution"]["duree"] and p.deroule == "messages_cles" and p.niveau == "allege"
    assert p.contenus == ["bivarie"]
    q = PresentationSpec.from_dict({"genre": "soutenance", "niveau": "allege"})
    assert q.niveau == "allege" and q.deroule == "classique"


def _executer(dataset, tmp_path, pres: dict, **cfg_kw):
    from analyste.config import Settings
    from analyste.pipeline import AnalysisConfig, run
    from analyste.writing.composer import RequestSpec
    cfg = AnalysisConfig(outcome="contraception_moderne",
                         explanatory=["groupe_age", "instruction", "milieu", "religion", "parite"],
                         cluster="grappe", redaction={"unite": "femmes"}, factorial=False, **cfg_kw)
    spec = RequestSpec(doc_type="memoire", title="Facteurs de l'utilisation de la contraception moderne",
                       author="Awa Ouédraogo", objectives=["Identifier les facteurs associés"],
                       context="La contraception moderne reste peu utilisée — surtout en milieu rural.",
                       problematique="Quels facteurs sont associés à l'utilisation ?", presentation=pres)
    return run(dataset, cfg, spec, tmp_path / "sortie", Settings(), b"k" * 32)


def _diapos(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for s in prs.slides:
        textes = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        for sh in s.shapes:
            if sh.has_table:
                textes += [c.text for r in sh.table.rows for c in r.cells]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        out.append(("\n".join(textes), notes))
    return prs, out


def test_soutenance_complete(dataset, tmp_path):
    res = _executer(dataset, tmp_path, {"active": True, "genre": "soutenance", "theme": "foret",
                                        "recommandations": ["Renforcer l'offre de services en milieu rural"]})
    assert "pptx" in res and zipfile.is_zipfile(res["pptx"])
    prs, diapos = _diapos(res["pptx"])
    tout = "\n".join(t + "\n" + n for t, n in diapos)
    assert not TIRETS.search(tout), TIRETS.search(tout)
    assert 6 <= len(diapos) <= round(20 / 1.5) + 1
    assert "Soutenance de mémoire" in diapos[0][0] and "Awa Ouédraogo" in diapos[0][0]
    assert diapos[0][1].startswith("Monsieur le Président du jury")
    assert "Renforcer l'offre de services en milieu rural" in tout
    assert "fois plus de chances" in tout
    assert all(n for _, n in diapos), "chaque diapositive porte des notes de l'orateur"
    assert prs.core_properties.author == "Awa Ouédraogo"
    with zipfile.ZipFile(res["pptx"]) as z:
        assert b"python-pptx" not in z.read("docProps/core.xml")
    with zipfile.ZipFile(res["docx"]) as z:
        assert b"python-docx" not in z.read("docProps/core.xml")


def test_restitution_messages_cles_et_annexes(dataset, tmp_path):
    res = _executer(dataset, tmp_path, {"active": True, "genre": "restitution", "format": "4:3", "notes": False})
    prs, diapos = _diapos(res["pptx"])
    assert prs.slide_width < prs.slide_height * 1.4  # format 4:3
    assert diapos[1][0].startswith("Ce qu'il faut retenir")
    assert any(t.startswith("Annexes") for t, _ in diapos)
    assert not any(n for _, n in diapos)
    # niveau allégé : pas de probabilité critique dans les diapositives de résultats principales
    principales = [t for t, _ in diapos[: [t for t, _ in diapos].index(next(t for t, _ in diapos
                                                                              if t.startswith("Annexes")))]]
    assert not any("p < 0,001" in t for t in principales)


def test_presentation_desactivee(dataset, tmp_path):
    res = _executer(dataset, tmp_path, {})
    assert "pptx" not in res


@pytest.mark.parametrize("correction,colonne", [("holm", "p (Holm)"), ("aucune", None)])
def test_correction_choisie(dataset, tmp_path, correction, colonne):
    from analyste.stats import bivariate
    sec = bivariate.bivariate(dataset, "contraception_moderne", ["instruction", "milieu", "religion"], tmp_path,
                              correction=correction)
    cols = list(sec.tables[0].data.columns)
    if colonne:
        assert colonne in cols and "Holm" in sec.paragraphs[0]
    else:
        assert not any(c.startswith("p (") for c in cols)
        assert "ne sont pas corrigées" in sec.paragraphs[0]


def test_seuil_et_tests_non_parametriques(dataset, tmp_path):
    from analyste.stats import bivariate, seuil
    jeton = seuil.definir(0.01)
    try:
        sec = bivariate.bivariate(dataset, "contraception_moderne", ["age", "parite", "instruction"], tmp_path,
                                  non_parametrique=True)
    finally:
        seuil.retablir(jeton)
    assert "au seuil de 1 %" in sec.paragraphs[0]
    tests = {r.test for r in sec.extra["tests"]}
    assert "t de Welch" not in tests and ("U de Mann-Whitney" in tests or "Kruskal-Wallis" in tests)
    assert seuil.alpha() == 0.05  # le seuil est rétabli après l'analyse


def test_analyses_choisies(dataset, tmp_path):
    from analyste.config import Settings
    from analyste.pipeline import AnalysisConfig, run
    from analyste.writing.composer import RequestSpec
    cfg = AnalysisConfig(outcome="contraception_moderne", explanatory=["instruction", "milieu"], cluster="grappe",
                         analyses=["bivarie"], factorial=False)
    res = run(dataset, cfg, RequestSpec(doc_type="article", title="Essai"), tmp_path / "s", Settings(), b"k" * 32)
    from docx import Document
    texte = "\n".join(p.text for p in Document(res["docx"]).paragraphs)
    assert "Analyse bivariée" in texte or "Facteurs associés" in texte
    assert "Analyse explicative multivariée" not in texte and "Analyse multi-niveaux" not in texte
