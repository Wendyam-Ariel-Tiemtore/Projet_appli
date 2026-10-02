"""Rédaction : règles de style (pas de tirets longs, nombres en lettres, articles, formulations des effets)."""

import re

from analyste.writing.phrases import Redac, accord, pcrit
from analyste.writing.style import (
    a,
    avec_article,
    de,
    enumeration,
    genre,
    lettres,
    nettoyer,
    nombre,
    ordinal,
    titre,
)

TIRETS = re.compile(r"[—–]")


def test_nombres_en_lettres():
    assert lettres(7) == "sept"
    assert lettres(21) == "vingt-et-un"
    assert lettres(71) == "soixante-et-onze"
    assert lettres(80) == "quatre-vingts"
    assert lettres(99) == "quatre-vingt-dix-neuf"
    assert nombre(7, "variables") == "sept (07) variables"
    assert nombre(1, "variables") == "une (01) variable"
    assert nombre(2511, "observations") == "2\u202f511 observations"  # espace fine ins\u00e9cable


def test_articles_et_contractions():
    assert avec_article("Niveau d'instruction") == "le niveau d'instruction"
    assert avec_article("Âge") == "l'âge"
    assert avec_article("Religion") == "la religion"
    assert avec_article("IMC") == "l'IMC"
    assert avec_article("Présence d'un CSPS") == "la présence d'un CSPS"
    assert de("le niveau d'instruction") == "du niveau d'instruction"
    assert de("les revenus") == "des revenus"
    assert a("le milieu de résidence") == "au milieu de résidence"
    assert a("l'exposition aux médias") == "à l'exposition aux médias"
    assert genre("parité") == "f" and genre("milieu de résidence") == "m"


def test_accords():
    assert accord("associé", "la religion") == "associée"
    assert accord("associé", "le niveau d'instruction") == "associé"
    assert accord("associé", "l'exposition aux médias") == "associée"
    assert enumeration(["a", "b", "c"]) == "a, b et c"
    assert ordinal(1, True) == "première" and ordinal(2) == "deuxième"


def test_nettoyage_des_tirets():
    assert nettoyer("Le modèle — estimé sur 200 cas — converge.") == "Le modèle (estimé sur 200 cas) converge."
    assert nettoyer("Un résultat net — les femmes instruites utilisent davantage.") == (
        "Un résultat net : les femmes instruites utilisent davantage.")
    assert not TIRETS.search(nettoyer("A – B — C"))
    assert titre("Chapitre 1 — Méthodologie") == "Chapitre 1 : Méthodologie"


def test_chances_et_probabilite_critique():
    assert Redac.chances(1.62) == "1,62 fois plus de chances"
    assert Redac.chances(0.74) == "26 % moins de chances"
    assert pcrit(0.0001) == "une probabilité critique inférieure à 0,001"
    assert pcrit(0.032, corrigee=True) == "une probabilité critique corrigée de 0,032"


def _all_text(sec):
    out = list(sec.paragraphs) + list(sec.warnings) + list(sec.method_notes) + list(sec.key_points)
    for t in sec.tables:
        out += [t.title, t.note] + [str(x) for x in t.data.to_numpy().ravel()]
    out += [f.caption for f in sec.figures]
    return "\n".join(out)


def test_textes_sans_tirets_et_au_style_memoire(dataset, tmp_path):
    from analyste.stats import bivariate, descriptive, models
    dataset.redaction = {"unite": "femmes", "evenement": "utiliser une méthode contraceptive moderne"}
    expl = ["groupe_age", "instruction", "milieu", "religion", "parite"]
    secs = [descriptive.describe(dataset, ["contraception_moderne"] + expl, tmp_path, outcome="contraception_moderne"),
            bivariate.bivariate(dataset, "contraception_moderne", expl, tmp_path),
            models.explain(dataset, "contraception_moderne", expl, tmp_path)]
    texte = "\n".join(_all_text(s) for s in secs)
    assert not TIRETS.search(texte), TIRETS.search(texte)
    assert "fois plus de chances d'utiliser une méthode contraceptive moderne" in texte
    assert "toutes choses égales par ailleurs" in texte.lower()
    assert "chez les femmes dont" in texte
    assert "« Niveau d'instruction »" not in texte  # les variables ne sont pas entre guillemets dans le texte


def test_document_complet_sans_tirets(dataset, tmp_path):
    from docx import Document as Docx

    from analyste.config import Settings
    from analyste.pipeline import AnalysisConfig, run
    from analyste.writing.composer import RequestSpec
    cfg = AnalysisConfig(outcome="contraception_moderne", explanatory=["instruction", "milieu", "religion"],
                         redaction={"unite": "femmes"}, factorial=False)
    spec = RequestSpec(doc_type="memoire", title="Essai — titre", objectives=["Identifier les facteurs"],
                       context="Contexte — saisi par l'auteur.")
    res = run(dataset, cfg, spec, tmp_path / "sortie", Settings(), b"k" * 32)
    d = Docx(res["docx"])
    texts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    joined = "\n".join(texts)
    assert not TIRETS.search(joined)
    assert "Chapitre 2 : Méthodologie" in joined
    assert "2.1. Source et nature des données" in joined  # sections numérotées dans leur chapitre
    assert "Tableau 2 : " in joined  # numéros de légende déjà calculés, sans mise à jour des champs


def test_formulation_personnalisee_appliquee(dataset):
    from analyste.pipeline import AnalysisConfig, apply_overrides
    cfg = AnalysisConfig(textes={"instruction": "le niveau d'instruction de la femme"},
                         redaction={"unite": "femmes", "indicateur": "la prévalence contraceptive"})
    apply_overrides(dataset, cfg)
    R = Redac(dataset, "contraception_moderne", "Oui")
    assert R.v("instruction") == "le niveau d'instruction de la femme"
    assert R.indicateur() == "la prévalence contraceptive"
    assert R.groupe("instruction", "Primaire") == ("les femmes dont le niveau d'instruction de la femme est "
                                                   "«\u00a0Primaire\u00a0»")
