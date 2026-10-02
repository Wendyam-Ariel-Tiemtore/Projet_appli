"""Génération du document Word (.docx) au format académique."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..stats.results import Figure, Table
from ..writing.composer import DOC_TYPES, Document, Item

FONT = "Times New Roman"


def _field(paragraph, instr: str, placeholder: str = "") -> None:
    """Insère un champ Word (TOC, PAGE, SEQ…)."""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run._r.append(b)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run._r.append(it)
    run = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run._r.append(s)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run._r.append(e)


def _styles(doc, line_spacing: float) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sizes = {1: 16, 2: 13, 3: 12, 4: 12}
    for lvl, size in sizes.items():
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = lvl == 4
        h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        rf = h.element.rPr.rFonts
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), FONT)
        h.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name in ("Caption",):
        if name in [s.name for s in doc.styles]:
            c = doc.styles[name]
            c.font.name = FONT
            c.font.size = Pt(10)
            c.font.bold = True
            c.font.italic = False
            c.font.color.rgb = RGBColor(0, 0, 0)


def _update_fields_on_open(doc) -> None:
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def _rich(paragraph, text: str, size: float | None = None, italic: bool = False) -> None:
    """Texte avec *italique* minimal (références APA)."""
    parts = re.split(r"(\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        else:
            r = paragraph.add_run(part)
            r.italic = italic
        if size:
            r.font.size = Pt(size)


def _caption(doc, label: str, title: str):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(f"{label} ")
    r.bold = True
    _field(p, f"SEQ {label} \\* ARABIC", "1")
    r2 = p.add_run(f" : {title}")
    r2.bold = True
    return p


def _set_cell_borders(cell, **kw) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, val in kw.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val.get("val", "single"))
        el.set(qn("w:sz"), str(val.get("sz", 8)))
        el.set(qn("w:color"), val.get("color", "000000"))
        borders.append(el)


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _no_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def add_table(doc, t: Table, landscape_hint: bool = False) -> None:
    _caption(doc, "Tableau", t.title)
    df = t.data
    ncol = len(df.columns)
    table = doc.add_table(rows=1, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _no_table_borders(table)
    widths = _column_widths(df)
    size = 9 if ncol <= 6 else 8 if ncol <= 9 else 7
    hdr = table.rows[0].cells
    for j, col in enumerate(df.columns):
        hdr[j].text = ""
        p = hdr[j].paragraphs[0]
        r = p.add_run(str(col))
        r.bold = True
        r.font.size = Pt(size)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        _shade(hdr[j], "E8ECF2")
        _set_cell_borders(hdr[j], top={"sz": 12}, bottom={"sz": 6})
    # En-tête répété sur chaque page
    trPr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    nrows = len(df)
    for i, (_, row) in enumerate(df.iterrows()):
        cells = table.add_row().cells
        is_group = all(str(v) == "" for v in list(row.values)[1:]) and str(row.values[0]) != ""
        for j, v in enumerate(row.values):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            r = p.add_run("" if v is None else str(v))
            r.font.size = Pt(size)
            if is_group and j == 0:
                r.bold = True
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if j <= 1 and not str(v)[:1].isdigit() \
                else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            if i == nrows - 1:
                _set_cell_borders(cells[j], bottom={"sz": 12})
    for row in table.rows:
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
    if t.note:
        p = doc.add_paragraph()
        _rich(p, "Note : " + t.note, size=9, italic=True)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    p = doc.add_paragraph()
    _rich(p, t.source, size=9, italic=True)
    p.paragraph_format.line_spacing = 1.0


def _column_widths(df, total_cm: float = 16.0) -> list:
    """Largeurs proportionnelles au contenu (racine de la longueur), bornées."""
    import math
    weights = []
    for col in df.columns:
        lens = [len(str(v)) for v in df[col].tolist()] + [len(str(col)) * 0.6]
        m = max(lens) if lens else 4
        avg = sum(lens) / len(lens) if lens else 4
        weights.append(max(2.0, math.sqrt(0.7 * m + 0.3 * avg)))
    tot = sum(weights)
    out = [max(1.2, total_cm * w / tot) for w in weights]
    scale = total_cm / sum(out)
    return [Cm(w * scale) for w in out]


def add_figure(doc, f: Figure) -> None:
    _caption(doc, "Graphique", f.caption)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(f.path), width=Cm(14.5))
    p = doc.add_paragraph()
    _rich(p, f.source, size=9, italic=True)


def _numbering_text(counters: list[int], level: int, text: str, numbered: bool) -> str:
    if not numbered:
        return text
    counters[level - 1] += 1
    for k in range(level, len(counters)):
        counters[k] = 0
    num = ".".join(str(c) for c in counters[:level] if c)
    return f"{num}. {text}" if level <= 3 else text


def build_docx(document: Document, out: Path) -> Path:
    spec = document.spec
    doc = DocxDocument()
    line = 1.15 if spec.doc_type == "article" else 1.5
    _styles(doc, line)
    _update_fields_on_open(doc)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, "PAGE", "1")

    # --- Page de garde ---
    for txt, size, bold in ((spec.institution, 13, True), ("", 12, False)):
        if txt:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt)
            r.font.size, r.bold = Pt(size), bold
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(DOC_TYPES.get(spec.doc_type, "Document").upper())
    r.font.size, r.bold = Pt(12), True
    r.font.color.rgb = RGBColor(0x55, 0x5F, 0x73)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(spec.title)
    r.font.size, r.bold = Pt(20), True
    r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    for _ in range(4):
        doc.add_paragraph()
    for lab, val in (("Présenté par", spec.author), ("Sous la direction de", spec.supervisor)):
        if val:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"{lab} : ")
            r.italic = True
            p.add_run(val).bold = True
    if spec.date_text:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(spec.date_text)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- Résumé ---
    if document.front.get("resume"):
        doc.add_heading("Résumé", level=1)
        for ptxt in document.front["resume"]:
            _para(doc, ptxt)
        if document.front.get("mots_cles"):
            p = doc.add_paragraph()
            p.add_run("Mots-clés : ").bold = True
            p.add_run(" ; ".join(document.front["mots_cles"]))
        if document.front.get("abstract"):
            doc.add_heading("Abstract", level=1)
            for ptxt in document.front["abstract"]:
                _para(doc, ptxt)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- Sommaire et listes ---
    doc.add_heading("Sommaire", level=1)
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "1-3" \\h \\z \\u', "Clic droit puis « Mettre à jour les champs » pour afficher le sommaire.")
    if spec.doc_type in ("memoire", "rapport_stage", "rapport_etude"):
        doc.add_heading("Liste des tableaux", level=1)
        _field(doc.add_paragraph(), 'TOC \\h \\z \\c "Tableau"', "Mettre à jour les champs pour afficher la liste.")
        doc.add_heading("Liste des graphiques", level=1)
        _field(doc.add_paragraph(), 'TOC \\h \\z \\c "Graphique"', "Mettre à jour les champs pour afficher la liste.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- Corps ---
    counters = [0, 0, 0, 0]
    numbered = spec.doc_type != "note_synthese"
    _render(doc, document.items, counters, numbered)

    # --- Références ---
    doc.add_heading("Références bibliographiques", level=1)
    for ref in document.references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-1.25)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        _rich(p, ref, size=11)

    # --- Annexes ---
    new = doc.add_section(WD_SECTION.NEW_PAGE)
    new.footer.is_linked_to_previous = True
    _render(doc, document.annex, [0, 0, 0, 0], False)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def _para(doc, text: str) -> None:
    p = doc.add_paragraph()
    _rich(p, text)


def _render(doc, items: list[Item], counters: list[int], numbered: bool) -> None:
    for it in items:
        if it.kind == "heading":
            doc.add_heading(_numbering_text(counters, it.level, it.text, numbered), level=min(it.level, 4))
        elif it.kind == "para":
            if it.text.strip().startswith("{") and it.text.strip().endswith("}"):
                p = doc.add_paragraph()
                r = p.add_run(it.text)
                r.font.name, r.font.size = "Consolas", Pt(8)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif it.text.strip():
                _para(doc, it.text)
        elif it.kind == "todo":
            p = doc.add_paragraph()
            r = p.add_run(it.text)
            r.italic = True
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif it.kind == "bullets":
            for b in it.items:
                p = doc.add_paragraph(style="List Bullet")
                _rich(p, b)
        elif it.kind == "table" and it.table is not None:
            add_table(doc, it.table)
        elif it.kind == "figure" and it.figure is not None:
            add_figure(doc, it.figure)
        elif it.kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def export_tables_xlsx(tables: list[Table], out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Sommaire"
    ws0.append(["N°", "Tableau"])
    for i, t in enumerate(tables, start=1):
        ws0.append([i, t.title])
        ws = wb.create_sheet(f"T{i}")
        ws.append([t.title])
        ws["A1"].font = Font(bold=True)
        ws.append([])
        ws.append([_safe_cell(c) for c in t.data.columns])
        for c in ws[3]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8ECF2")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in t.data.itertuples(index=False):
            ws.append([_safe_cell(v) for v in row])
        ws.append([])
        ws.append([_safe_cell("Note : " + t.note)])
        ws.append([_safe_cell(t.source)])
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in list(col)[2:]) if ws.max_row > 2 else 12
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.save(out)
    return out


def _safe_cell(v) -> str:
    """Neutralise l'injection de formules dans les tableurs (CSV/Excel injection)."""
    s = "" if v is None else str(v)
    if s[:1] in ("=", "+", "-", "@", "\t", "\r") and not re.match(r"^-?\d", s):
        return "'" + s
    return s
