"""Rendu PowerPoint d'un plan de présentation (writing/presentation.py).

Choix de mise en page : fonds sombres pour l'ouverture, la conclusion et la clôture, fonds blancs pour le contenu ;
motif unique de pastilles numérotées ; aucune ligne décorative sous les titres ni bandeau de couleur ; graphiques
natifs modifiables dans PowerPoint ; polices disponibles partout (Cambria pour les titres, Calibri pour le texte).
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ..stats import fmt
from ..writing.presentation import THEME_COULEURS, Diapo, Plan
from ..writing.style import nettoyer
from ..writing.style import titre as titre_propre

TITRE, CORPS = "Cambria", "Calibri"
H = 7.5
M = 0.6  # marge latérale

PALETTES = THEME_COULEURS


def typo(texte: str) -> str:
    """Espaces insécables du français : un nombre reste avec son unité, les guillemets avec leur contenu, et
    aucune ligne ne commence par un deux-points, un point-virgule ou un point d'interrogation."""
    t = re.sub(r"(\d) (%|ans\b|points\b)", "\\1\u00a0\\2", texte)
    t = t.replace("« ", "«\u00a0").replace(" »", "\u00a0»")
    return re.sub(r" ([:;?!])", "\u00a0\\1", t)


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


class Rendu:
    def __init__(self, plan: Plan):
        self.plan = plan
        self.c = PALETTES.get(plan.spec.theme, PALETTES["ardoise"])
        self.W = 13.333 if plan.spec.format == "16:9" else 10.0
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.W)
        self.prs.slide_height = Inches(H)
        self.vide = self.prs.slide_layouts[6]
        self.numero = 0

    # ------------------------------------------------------------------ outils
    def _fond(self, slide, couleur: str) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(couleur)

    def _boite(self, slide, x, y, w, h, nom: str = ""):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        if nom:
            tb.name = nom
        return tb, tf

    @staticmethod
    def _run(p, texte: str, taille: float, couleur: str, gras=False, italique=False, police=CORPS):
        r = p.add_run()
        r.text = typo(texte)
        f = r.font
        f.size, f.bold, f.italic, f.name = Pt(taille), gras, italique, police
        f.color.rgb = _rgb(couleur)
        return r

    def _texte(self, slide, x, y, w, h, texte, taille, couleur, gras=False, italique=False, police=CORPS,
               align=PP_ALIGN.LEFT, ancre=MSO_ANCHOR.TOP, nom=""):
        tb, tf = self._boite(slide, x, y, w, h, nom)
        tf.vertical_anchor = ancre
        p = tf.paragraphs[0]
        p.alignment = align
        p.line_spacing = 1.05
        self._run(p, texte, taille, couleur, gras, italique, police)
        return tb

    @staticmethod
    def _puce(p, couleur: str, retrait: float = 0.28) -> None:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(retrait))))
        pPr.set("indent", str(-int(Inches(retrait))))
        for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        buClr = pPr.makeelement(qn("a:buClr"), {})
        srgb = buClr.makeelement(qn("a:srgbClr"), {"val": couleur})
        buClr.append(srgb)
        pPr.append(buClr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))

    @staticmethod
    def ajuste(items: list[str], w: float, h: float, tailles=(18, 16, 15, 14), retrait: float = 0.28,
               espace: float = 8) -> tuple[int, list[str]]:
        """Plus grande taille de police pour laquelle les puces tiennent dans la boîte ; sinon, moins de puces."""
        items = [i for i in items if i]
        while True:
            for t in tailles:
                cpl = max(12, int((w - retrait) * 72 / (t * 0.53)))
                lignes = sum(max(1, math.ceil(len(s) / cpl)) for s in items)
                haut = (lignes * t * 1.18 + espace * max(0, len(items) - 1)) / 72
                if haut <= h:
                    return t, items
            if len(items) <= 1:
                return tailles[-1], items
            items = items[:-1]

    def _puces(self, slide, x, y, w, h, items, couleur=None, puce=None, tailles=(18, 16, 15, 14), nom="puces"):
        taille, items = self.ajuste(items, w, h, tailles)
        tb, tf = self._boite(slide, x, y, w, h, nom)
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.08
            p.space_after = Pt(8)
            self._puce(p, puce or self.c["accent"])
            self._run(p, it, taille, couleur or self.c["texte"])
        return tb

    def _taille_titre(self, texte: str) -> int:
        w = self.W - 2 * M
        for t in (30, 28, 26, 24):
            if math.ceil(len(texte) / int(w * 72 / (t * 0.56))) <= 2:
                return t
        return 22

    def _titre(self, slide, texte: str, couleur=None) -> None:
        """Même position et même taille de titre sur toutes les diapositives de contenu."""
        texte = titre_propre(texte).rstrip(" .")
        self._texte(slide, M, 0.42, self.W - 2 * M, 1.1, texte, self.taille_titre, couleur or self.c["fonce"],
                    gras=True, police=TITRE, ancre=MSO_ANCHOR.BOTTOM, nom="Titre")

    def _pied(self, slide, annexe: bool = False) -> None:
        self.numero += 1
        y = H - 0.45
        pied = self.plan.pied
        if annexe:
            pied = (pied + " | Annexe") if pied else "Annexe"
        if pied:
            self._texte(slide, M, y, self.W * 0.6, 0.3, pied, 10, self.c["attenue"], nom="Pied")
        self._texte(slide, self.W - M - 1.0, y, 1.0, 0.3, str(self.numero), 10, self.c["attenue"],
                    align=PP_ALIGN.RIGHT, nom="Numéro")

    def _carte(self, slide, x, y, w, h, couleur=None):
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        s.adjustments[0] = 0.06
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(couleur or self.c["doux"])
        s.line.fill.background()
        s.shadow.inherit = False
        s.name = "Carte"
        return s

    def _pastille(self, slide, x, y, d, texte, fond=None, couleur="FFFFFF", taille=14):
        s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(fond or self.c["accent"])
        s.line.fill.background()
        s.shadow.inherit = False
        tf = s.text_frame
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        self._run(p, texte, taille, couleur, gras=True)
        s.name = "Pastille"
        return s

    def _note_bas(self, slide, texte: str, x=None, w=None, y=None) -> None:
        if texte:
            self._texte(slide, M if x is None else x, (H - 0.95) if y is None else y, (self.W - 2 * M) if w is None
                        else w, 0.4, nettoyer(texte), 11, self.c["attenue"], italique=True, nom="Lecture")

    def _encadre(self, slide, x, y, w, h, etiquette: str, texte: str, taille_max=None) -> None:
        taille_max = taille_max or (24 if self.W > 11 else 20)
        self._carte(slide, x, y, w, h)
        self._texte(slide, x + 0.35, y + 0.3, w - 0.7, 0.35, etiquette, 12, self.c["accent"], gras=True)
        taille, _ = self.ajuste([texte], w - 0.7, h - 1.2, tailles=(taille_max, 22, 20, 18, 16, 15), retrait=0)
        self._texte(slide, x + 0.35, y + 0.8, w - 0.7, h - 1.2, texte, taille, self.c["fonce"], italique=True,
                    police=TITRE, ancre=MSO_ANCHOR.MIDDLE)

    def _commentaire(self, slide, x, y, w, h, puces: list[str], etiquette="Lecture") -> None:
        self._carte(slide, x, y, w, h)
        self._texte(slide, x + 0.3, y + 0.28, w - 0.6, 0.35, etiquette, 12, self.c["accent"], gras=True)
        self._puces(slide, x + 0.3, y + 0.75, w - 0.6, h - 1.0, puces, tailles=(16, 15, 14))

    @staticmethod
    def _notes(slide, texte: str) -> None:
        if texte:
            slide.notes_slide.notes_text_frame.text = texte

    # ------------------------------------------------------------------ diapositives
    def d_titre(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, self.c["fonce"])
        rayon = 2.3 if self.W > 11 else 1.6
        cx = self.W - M - 2 * rayon + 0.4
        cercle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(1.25), Inches(2 * rayon), Inches(2 * rayon))
        cercle.fill.solid()
        cercle.fill.fore_color.rgb = _rgb(self.c["relief"])
        cercle.line.fill.background()
        cercle.name = "Motif"
        petit = 0.55
        p2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.2), Inches(1.25 + 2 * rayon - petit - 0.1),
                                Inches(petit), Inches(petit))
        p2.fill.solid()
        p2.fill.fore_color.rgb = _rgb(self.c["accent"])
        p2.line.fill.background()
        p2.name = "Motif"
        w = cx - M - 0.4
        self._texte(s, M, 1.2, w, 0.4, d.sous_titre, 16, self.c["accent_clair"], gras=True)
        taille = 36
        for t in (36, 32, 28, 26, 24):
            cpl = int(w * 72 / (t * 0.56))
            if math.ceil(len(d.titre) / cpl) * t * 1.15 / 72 <= 2.7:
                taille = t
                break
        self._texte(s, M, 1.7, w, 2.8, d.titre, taille, "FFFFFF", gras=True, police=TITRE)
        lignes = []
        if d.meta.get("auteur"):
            lignes.append(("Présenté par ", d.meta["auteur"]))
        if d.meta.get("direction"):
            lignes.append(("Sous la direction de ", d.meta["direction"]))
        tb, tf = self._boite(s, M, 4.75, w, 1.3, "Auteurs")
        for i, (lab, val) in enumerate(lignes):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            self._run(p, lab, 15, self.c["accent_clair"])
            self._run(p, val, 15, "FFFFFF", gras=True)
        bas = " | ".join(x for x in (d.meta.get("institution"), d.meta.get("date")) if x)
        if bas:
            self._texte(s, M, H - 1.0, self.W - 2 * M, 0.4, bas, 13, "D5DCE2")
        self.numero += 1
        self._notes(s, d.notes)

    def d_plan(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        items = d.puces[:10]
        deux = len(items) > 5
        colw = (self.W - 2 * M - (0.6 if deux else 0)) / (2 if deux else 1)
        for i, it in enumerate(items):
            col, lig = (i // 5, i % 5) if deux else (0, i)
            x = M + col * (colw + 0.6)
            y = 1.95 + lig * 0.95
            self._pastille(s, x, y, 0.58, str(i + 1), taille=16)
            self._texte(s, x + 0.85, y, colw - 0.9, 0.58, it, 20, self.c["texte"], ancre=MSO_ANCHOR.MIDDLE)
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def _lignes_numerotees(self, s, items, couleur_texte, fond_pastille=None, y0=1.95, h_total=4.7, taille_max=20):
        n = max(1, len(items))
        pas = h_total / n
        w = self.W - 2 * M - 0.9
        taille = min(self.ajuste([it], w, pas - 0.2, tailles=(taille_max, 18, 16, 15, 14), retrait=0)[0]
                     for it in items) if items else 16
        for i, it in enumerate(items):
            y = y0 + i * pas
            self._pastille(s, M, y + 0.02, 0.55, str(i + 1), fond=fond_pastille, taille=15)
            self._texte(s, M + 0.85, y, w, pas - 0.15, it, taille, couleur_texte)

    def d_messages(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        self._lignes_numerotees(s, d.puces[:4], self.c["texte"])
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_conclusion(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, self.c["fonce"])
        self._titre(s, d.titre, couleur="FFFFFF")
        self._lignes_numerotees(s, d.puces[:3], "FFFFFF", fond_pastille=self.c["accent"])
        self.numero += 1
        self._notes(s, d.notes)

    def d_texte(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        if d.encadre:
            wg = (self.W - 2 * M) * 0.56
            self._puces(s, M, 2.0, wg, 4.4, d.puces, tailles=(22, 20, 18, 16, 15, 14))
            xe = M + wg + 0.5
            self._encadre(s, xe, 1.95, self.W - M - xe, 4.4, d.meta.get("etiquette", "À retenir"), d.encadre)
        else:
            self._lignes_numerotees(s, d.puces[:4], self.c["texte"])
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_cartes(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        cartes = d.cartes[:4]
        n = len(cartes)
        cols = 2 if n == 4 else max(1, n)
        rows = 2 if n == 4 else 1
        gap = 0.4
        cw = (self.W - 2 * M - gap * (cols - 1)) / cols
        ch = (4.6 - gap * (rows - 1)) / rows
        if rows == 1:
            cpl = max(12, int((cw - 0.6) * 72 / (20 * 0.53)))  # caractères par ligne à 20 points
            besoin = max(math.ceil(len(t) / cpl) for _, t in cartes)
            ch = min(ch, max(2.4, 1.45 + besoin * 20 * 1.2 / 72))
        for i, (lab, txt) in enumerate(cartes):
            r, c = divmod(i, cols)
            x, y = M + c * (cw + gap), 1.95 + r * (ch + gap)
            self._carte(s, x, y, cw, ch)
            self._pastille(s, x + 0.3, y + 0.3, 0.5, str(i + 1), taille=14)
            self._texte(s, x + 0.95, y + 0.3, cw - 1.2, 0.5, lab, 14, self.c["primaire"], gras=True,
                        ancre=MSO_ANCHOR.MIDDLE)
            taille, _ = self.ajuste([txt], cw - 0.6, ch - 1.2, tailles=(20, 18, 16, 15, 14), retrait=0)
            self._texte(s, x + 0.3, y + 1.05, cw - 0.6, ch - 1.25, txt, taille, self.c["texte"])
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_chiffres(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        chiffres = d.chiffres[:4]
        n = max(1, len(chiffres))
        gap = 0.35
        if not d.puces and n >= 3:  # grille 2 × 2, chiffres plus grands
            cw = (self.W - 2 * M - gap) / 2
            ch = 1.95
            for i, (val, lab) in enumerate(chiffres):
                r, c = divmod(i, 2)
                x, y = M + c * (cw + gap), 1.95 + r * (ch + gap)
                self._carte(s, x, y, cw, ch)
                self._texte(s, x + 0.4, y + 0.3, cw - 0.8, 0.9, val, 44, self.c["accent"] if i == 0 else
                            self.c["primaire"], gras=True, police=TITRE)
                self._texte(s, x + 0.4, y + 1.2, cw - 0.8, 0.6, lab, 16, self.c["texte"])
            y = 1.95 + 2 * ch + gap
        else:
            cw = (self.W - 2 * M - gap * (n - 1)) / n
            haut = 1.75
            for i, (val, lab) in enumerate(chiffres):
                x = M + i * (cw + gap)
                self._carte(s, x, 1.95, cw, haut)
                tv = 40 if len(val) <= 7 else 32
                self._texte(s, x + 0.25, 2.1, cw - 0.5, 0.8, val, tv, self.c["accent"] if i == 0 else
                            self.c["primaire"], gras=True, police=TITRE)
                self._texte(s, x + 0.25, 2.95, cw - 0.5, haut - 1.1, lab, 14, self.c["attenue"])
            y = 1.95 + haut + 0.35
        if d.puces:
            self._texte(s, M, y, 4.0, 0.35, "Méthodes d'analyse" if d.titre.startswith("Données") else "À retenir",
                        12, self.c["accent"], gras=True)
            self._puces(s, M, y + 0.45, self.W - 2 * M, H - y - 1.6, d.puces, tailles=(16, 15, 14))
        if d.encadre:
            self._note_bas(s, d.encadre)
        elif d.note_bas:
            self._note_bas(s, d.note_bas)
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_graphique(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        g = d.graphique or {}
        wg = (self.W - 2 * M) * 0.58
        cats = [str(c) for c in g.get("categories", [])]
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series(g.get("serie", "Valeur"), [round(float(v), 2) for v in g.get("valeurs", [])])
        horiz = len(cats) > 5 or max((len(c) for c in cats), default=0) > 18
        typ = XL_CHART_TYPE.BAR_CLUSTERED if horiz else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = s.shapes.add_chart(typ, Inches(M), Inches(1.9), Inches(wg), Inches(4.3), cd)
        gf.name = "Graphique"
        ch = gf.chart
        ch.has_legend = False
        ch.has_title = False
        ch.font.size, ch.font.name = Pt(12), CORPS
        ch.font.color.rgb = _rgb(self.c["texte"])
        plot = ch.plots[0]
        plot.gap_width = 70
        plot.vary_by_categories = False
        ser = plot.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb(self.c["primaire"])
        hi = g.get("surligne")
        if hi is not None and 0 <= hi < len(cats):
            pt = ser.points[hi]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(self.c["accent"])
        # étiquettes écrites à la française (virgule décimale) quelle que soit la langue de PowerPoint
        for i, v in enumerate(g.get("valeurs", [])):
            dlab = ser.points[i].data_label
            dlab.position = XL_LABEL_POSITION.OUTSIDE_END
            tf = dlab.text_frame
            tf.text = typo(fmt.pct(float(v) / 100) if g.get("format") == "pct" else fmt.num(v))
            for par in tf.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(12)
                    run.font.name = CORPS
                    run.font.color.rgb = _rgb(self.c["texte"])
        va = ch.value_axis
        va.has_major_gridlines = False
        va.visible = False
        vmax = max([float(v) for v in g.get("valeurs", [1])] + [1])
        va.maximum_scale = vmax * 1.18
        va.minimum_scale = 0
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(12)
        ca.tick_labels.font.color.rgb = _rgb(self.c["texte"])
        ca.format.line.color.rgb = _rgb("C9D1D6")
        ca.has_major_gridlines = False
        if horiz:
            ca.reverse_order = True
        self._note_bas(s, d.note_bas, w=wg, y=6.25)
        xe = M + wg + 0.45
        self._commentaire(s, xe, 1.95, self.W - M - xe, 4.3, d.puces)
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_image(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        wg = (self.W - 2 * M) * (0.58 if d.puces else 1.0)
        hg = 4.3
        if d.image and Path(d.image).exists():
            from PIL import Image
            with Image.open(d.image) as im:
                ratio = im.width / im.height
            w, h = (wg, wg / ratio) if wg / ratio <= hg else (hg * ratio, hg)
            pic = s.shapes.add_picture(str(d.image), Inches(M + (wg - w) / 2), Inches(1.9 + (hg - h) / 2),
                                       Inches(w), Inches(h))
            pic.name = "Figure"
        self._note_bas(s, d.note_bas, w=wg, y=6.25)
        if d.puces:
            xe = M + wg + 0.45
            self._commentaire(s, xe, 1.95, self.W - M - xe, 4.3, d.puces)
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_tableau(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, "FFFFFF")
        self._titre(s, d.titre)
        df = d.tableau
        # tableau large ou format 4:3 : pleine largeur, le commentaire reste dans les notes de l'orateur
        avec_puces = bool(d.puces) and self.W > 11 and (df is None or len(df.columns) <= 4)
        wt = (self.W - 2 * M) * (0.64 if avec_puces else 1.0)
        if df is not None and len(df.columns):
            df = df.fillna("").astype(str)
            nr, nc = len(df) + 1, len(df.columns)
            taille = 13 if nr <= 8 else 12
            rh = 0.42 if nr <= 8 else 0.37
            longueurs = [max([len(str(c))] + [len(v) for v in df[c]]) for c in df.columns]
            tot = sum(min(max(lg, 6), 40) for lg in longueurs)
            largeurs = [wt * min(max(lg, 6), 40) / tot for lg in longueurs]
            gf = s.shapes.add_table(nr, nc, Inches(M), Inches(1.95), Inches(wt), Inches(rh * nr))
            gf.name = "Tableau"
            tbl = gf.table
            tbl.first_row = True
            tbl.horz_banding = False
            for j, w in enumerate(largeurs):
                tbl.columns[j].width = Emu(int(Inches(w)))
            for i in range(nr):
                tbl.rows[i].height = Emu(int(Inches(rh)))
                for j in range(nc):
                    cell = tbl.cell(i, j)
                    txt = str(df.columns[j]) if i == 0 else df.iat[i - 1, j]
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(self.c["primaire"] if i == 0 else
                                                    ("FFFFFF" if i % 2 else self.c["doux"]))
                    cell.margin_left = cell.margin_right = Inches(0.08)
                    cell.margin_top = cell.margin_bottom = Inches(0.03)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    tf = cell.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.LEFT if j == 0 or i == 0 else PP_ALIGN.LEFT
                    self._run(p, txt, taille, "FFFFFF" if i == 0 else self.c["texte"], gras=(i == 0))
            y_note = min(1.95 + rh * nr + 0.2, H - 0.95)
            self._note_bas(s, d.note_bas, w=wt, y=y_note)
        if avec_puces:
            xe = M + wt + 0.45
            self._commentaire(s, xe, 1.95, self.W - M - xe, 4.3, d.puces, etiquette="À retenir")
        self._pied(s, d.annexe)
        self._notes(s, d.notes)

    def d_fin(self, d: Diapo) -> None:
        s = self.prs.slides.add_slide(self.vide)
        self._fond(s, self.c["fonce"])
        self._pastille(s, M, 2.2, 0.7, "", fond=self.c["accent"])
        self._texte(s, M, 3.1, self.W - 2 * M, 1.0, d.titre, 40, "FFFFFF", gras=True, police=TITRE)
        self._texte(s, M, 4.15, self.W - 2 * M, 0.8, d.sous_titre, 20, self.c["accent_clair"])
        bas = " | ".join(x for x in (self.plan.pied, d.meta.get("contact")) if x)
        if bas:
            self._texte(s, M, H - 1.1, self.W - 2 * M, 0.4, bas, 14, "D5DCE2")
        self.numero += 1
        self._notes(s, d.notes)

    # ------------------------------------------------------------------ assemblage
    def construire(self, chemin: Path) -> Path:
        contenus = [d.titre for d in self.plan.diapos if d.kind not in ("titre", "fin") and d.titre]
        self.taille_titre = min([self._taille_titre(t) for t in contenus] or [30])
        for d in self.plan.diapos:
            getattr(self, f"d_{d.kind}", self.d_texte)(d)
        from .docx_builder import _proprietes
        _proprietes(self.prs.core_properties, self.plan.titre, self.plan.pied)
        chemin.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(chemin)
        return chemin


def build_pptx(plan: Plan, out: Path) -> Path:
    return Rendu(plan).construire(out)
